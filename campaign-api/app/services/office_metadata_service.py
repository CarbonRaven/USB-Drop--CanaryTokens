"""Office document metadata injection service."""

from docx import Document
from openpyxl import load_workbook
from datetime import datetime, timedelta
from io import BytesIO
import random
import zipfile
import xml.etree.ElementTree as ET
from typing import Optional, Dict, Any
import logging

logger = logging.getLogger(__name__)


class OfficeMetadataService:
    """Inject realistic metadata into Office documents."""

    # Common author names for documents
    AUTHORS = [
        "Jennifer Williams",
        "Michael Chen",
        "Sarah Johnson",
        "David Kim",
        "Emily Rodriguez",
        "James Wilson",
        "Amanda Thompson",
        "Robert Martinez",
        "Lisa Anderson",
        "Christopher Lee",
        "Jessica Brown",
        "Daniel Taylor",
        "Michelle Davis",
        "Kevin White",
        "Stephanie Garcia",
    ]

    # Company names by scenario type
    COMPANIES = {
        "default": [
            "Acme Corporation",
            "GlobalTech Industries",
            "Innovative Solutions Inc",
            "Premier Services LLC",
            "United Enterprises",
        ],
        "corporate": [
            "Corporate Holdings Inc",
            "Enterprise Solutions Group",
            "Strategic Partners LLC",
            "Executive Services Corp",
        ],
        "finance": [
            "Financial Services Corp",
            "Capital Management LLC",
            "Investment Partners Inc",
            "Fiscal Solutions Group",
        ],
        "hr": [
            "Human Resources Dept",
            "People Operations",
            "Talent Management",
            "HR Services",
        ],
        "it_department": [
            "IT Department",
            "Technology Services",
            "Information Systems",
            "Tech Operations",
        ],
    }

    # Department-specific author pools
    DEPARTMENT_AUTHORS = {
        "hr": [
            "Sandra Mitchell - HR",
            "Thomas Wright - HR Director",
            "Karen Phillips - Benefits",
            "Richard Moore - Recruiting",
        ],
        "finance": [
            "Patricia Clark - Finance",
            "Steven Hall - Controller",
            "Nancy Allen - Accounting",
            "George Young - CFO Office",
        ],
        "it_department": [
            "Brian Scott - IT",
            "Melissa Green - Systems Admin",
            "Jeffrey Adams - Network Ops",
            "Laura Nelson - IT Security",
        ],
        "executive": [
            "John Smith - CEO Office",
            "Elizabeth Carter - Executive Admin",
            "William Turner - Strategy",
            "Margaret Hill - Board Secretary",
        ],
    }

    @classmethod
    def get_authors_for_scenario(cls, scenario_type: Optional[str] = None) -> list:
        """Get appropriate author pool for scenario."""
        if scenario_type and scenario_type in cls.DEPARTMENT_AUTHORS:
            return cls.DEPARTMENT_AUTHORS[scenario_type] + cls.AUTHORS[:5]
        return cls.AUTHORS

    @classmethod
    def get_company_for_scenario(cls, scenario_type: Optional[str] = None) -> str:
        """Get appropriate company name for scenario."""
        if scenario_type and scenario_type in cls.COMPANIES:
            return random.choice(cls.COMPANIES[scenario_type])
        return random.choice(cls.COMPANIES["default"])

    @classmethod
    def enhance_docx(
        cls,
        doc_data: bytes,
        author: Optional[str] = None,
        company: Optional[str] = None,
        created: Optional[datetime] = None,
        modified: Optional[datetime] = None,
        editing_minutes: Optional[int] = None,
        title: Optional[str] = None,
        scenario_type: Optional[str] = None,
    ) -> bytes:
        """
        Enhance Word document with realistic metadata.

        Args:
            doc_data: Raw document bytes
            author: Document author (random if None)
            company: Company name (random if None)
            created: Creation date (random if None)
            modified: Modification date (random if None)
            editing_minutes: Total editing time (calculated if None)
            title: Document title
            scenario_type: Scenario for context-appropriate metadata

        Returns:
            Enhanced document bytes
        """
        try:
            doc = Document(BytesIO(doc_data))

            # Set defaults based on scenario
            authors = cls.get_authors_for_scenario(scenario_type)
            if author is None:
                author = random.choice(authors)
            if company is None:
                company = cls.get_company_for_scenario(scenario_type)
            if modified is None:
                modified = datetime.now() - timedelta(days=random.randint(1, 30))
            if created is None:
                created = modified - timedelta(days=random.randint(30, 180))

            # Estimate editing time based on content if not provided
            if editing_minutes is None:
                para_count = len(doc.paragraphs)
                page_estimate = max(1, para_count // 15)
                editing_minutes = cls._calculate_editing_time(page_estimate)

            # Update core properties
            core_props = doc.core_properties
            core_props.author = author
            core_props.last_modified_by = random.choice([author] + authors[:3])
            core_props.created = created
            core_props.modified = modified
            if title:
                core_props.title = title

            # Save to bytes
            output = BytesIO()
            doc.save(output)
            doc_bytes = output.getvalue()

            # Update app.xml for editing time and company
            doc_bytes = cls._update_app_xml(doc_bytes, company, editing_minutes)

            logger.debug(f"Enhanced DOCX: author={author}, company={company}, time={editing_minutes}min")
            return doc_bytes

        except Exception as e:
            logger.error(f"Failed to enhance DOCX metadata: {e}")
            return doc_data  # Return original on failure

    @classmethod
    def enhance_xlsx(
        cls,
        xlsx_data: bytes,
        author: Optional[str] = None,
        company: Optional[str] = None,
        created: Optional[datetime] = None,
        modified: Optional[datetime] = None,
        title: Optional[str] = None,
        scenario_type: Optional[str] = None,
    ) -> bytes:
        """
        Enhance Excel document with realistic metadata.

        Args:
            xlsx_data: Raw spreadsheet bytes
            author: Document author (random if None)
            company: Company name (random if None)
            created: Creation date (random if None)
            modified: Modification date (random if None)
            title: Document title
            scenario_type: Scenario for context-appropriate metadata

        Returns:
            Enhanced spreadsheet bytes
        """
        try:
            wb = load_workbook(BytesIO(xlsx_data))

            # Set defaults based on scenario
            authors = cls.get_authors_for_scenario(scenario_type)
            if author is None:
                author = random.choice(authors)
            if company is None:
                company = cls.get_company_for_scenario(scenario_type)
            if modified is None:
                modified = datetime.now() - timedelta(days=random.randint(1, 30))
            if created is None:
                created = modified - timedelta(days=random.randint(30, 180))

            # Update properties
            wb.properties.creator = author
            wb.properties.lastModifiedBy = random.choice([author] + authors[:3])
            wb.properties.created = created
            wb.properties.modified = modified
            if title:
                wb.properties.title = title
            if company:
                wb.properties.company = company

            # Save to bytes
            output = BytesIO()
            wb.save(output)

            logger.debug(f"Enhanced XLSX: author={author}, company={company}")
            return output.getvalue()

        except Exception as e:
            logger.error(f"Failed to enhance XLSX metadata: {e}")
            return xlsx_data  # Return original on failure

    @classmethod
    def _calculate_editing_time(cls, page_count: int) -> int:
        """
        Calculate realistic editing time in minutes.

        Simulates time spent writing/editing a document based on page count.
        """
        # Base time per page (15-30 minutes)
        base_per_page = random.randint(15, 30)
        # Initial setup/overhead time
        overhead = random.randint(10, 30)
        # Random variance factor
        variance = random.uniform(0.8, 1.2)
        # Multiple editing sessions add time
        revision_multiplier = random.uniform(1.0, 1.5)

        total = (page_count * base_per_page + overhead) * variance * revision_multiplier
        return int(total)

    @classmethod
    def _update_app_xml(
        cls,
        doc_bytes: bytes,
        company: str,
        editing_minutes: int
    ) -> bytes:
        """
        Update app.xml inside docx for TotalTime, Company, and Application info.

        Word documents (docx) are ZIP archives containing XML files.
        The app.xml file contains extended properties like editing time.
        """
        input_zip = BytesIO(doc_bytes)
        output_zip = BytesIO()

        try:
            with zipfile.ZipFile(input_zip, 'r') as zin:
                with zipfile.ZipFile(output_zip, 'w', zipfile.ZIP_DEFLATED) as zout:
                    for item in zin.infolist():
                        data = zin.read(item.filename)

                        if item.filename == 'docProps/app.xml':
                            data = cls._modify_app_xml(data, company, editing_minutes)

                        zout.writestr(item, data)

            return output_zip.getvalue()

        except Exception as e:
            logger.error(f"Failed to update app.xml: {e}")
            return doc_bytes

    @classmethod
    def _modify_app_xml(cls, xml_data: bytes, company: str, editing_minutes: int) -> bytes:
        """Modify the app.xml content with new metadata."""
        ns = {'ep': 'http://schemas.openxmlformats.org/officeDocument/2006/extended-properties'}

        # Register namespace to avoid ns0 prefix
        ET.register_namespace('', 'http://schemas.openxmlformats.org/officeDocument/2006/extended-properties')
        ET.register_namespace('vt', 'http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes')

        root = ET.fromstring(xml_data)

        # Helper to find or create element
        def find_or_create(tag_name: str) -> ET.Element:
            elem = root.find(f'ep:{tag_name}', ns)
            if elem is None:
                elem = ET.SubElement(
                    root,
                    f'{{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}}{tag_name}'
                )
            return elem

        # Update TotalTime
        find_or_create('TotalTime').text = str(editing_minutes)

        # Update Company
        find_or_create('Company').text = company

        # Update Application to look like real MS Word
        find_or_create('Application').text = 'Microsoft Office Word'

        # Update AppVersion to current Office version
        app_versions = ['16.0000', '15.0000', '14.0000']  # Office 2016+, 2013, 2010
        find_or_create('AppVersion').text = random.choice(app_versions)

        return ET.tostring(root, encoding='UTF-8', xml_declaration=True)


def enhance_office_document(
    doc_data: bytes,
    doc_type: str,
    scenario_type: Optional[str] = None,
    file_timestamp: Optional[datetime] = None,
    **kwargs
) -> bytes:
    """
    Convenience function to enhance any Office document.

    Args:
        doc_data: Raw document bytes
        doc_type: Type of document ('ms_word' or 'ms_excel')
        scenario_type: Scenario for context-appropriate metadata
        file_timestamp: Use this as the modified date
        **kwargs: Additional arguments passed to enhance methods

    Returns:
        Enhanced document bytes
    """
    if file_timestamp:
        kwargs['modified'] = file_timestamp

    if doc_type == 'ms_word':
        return OfficeMetadataService.enhance_docx(
            doc_data,
            scenario_type=scenario_type,
            **kwargs
        )
    elif doc_type == 'ms_excel':
        return OfficeMetadataService.enhance_xlsx(
            doc_data,
            scenario_type=scenario_type,
            **kwargs
        )
    else:
        logger.warning(f"Unknown document type: {doc_type}")
        return doc_data


def get_available_authors() -> list:
    """Get list of all available author names."""
    return OfficeMetadataService.AUTHORS


def get_available_companies() -> Dict[str, list]:
    """Get all company names organized by scenario."""
    return OfficeMetadataService.COMPANIES
